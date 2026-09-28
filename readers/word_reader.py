"""Reader for Microsoft Word documents."""

from __future__ import annotations

from pathlib import Path

from docx import Document as DocxDocument

from models.document import Document
from readers.base_reader import BaseReader
from services.exceptions import DocumentReadError, InvalidDocumentError


class WordReader(BaseReader):
    """Read .docx files into a common Document model."""

    def can_read(self, source: str) -> bool:
        return Path(source).suffix.lower() == ".docx"

    def read(self, source: str) -> Document:
        path = Path(source)
        if not path.exists():
            raise InvalidDocumentError(
                "Unable to read this document.\n\nPlease verify that the file is valid and try again."
            )

        try:
            document = DocxDocument(str(path))
        except Exception as exc:
            raise InvalidDocumentError(
                "Unable to read this document.\n\nPlease verify that the file is valid and try again."
            ) from exc

        sections: list[str] = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                sections.append(text)

        for table in document.tables:
            rows = []
            for row in table.rows:
                rows.append(" | ".join(cell.text.strip() for cell in row.cells))
            if rows:
                sections.append("\n".join(rows))

        content = "\n\n".join(sections).strip()

        return Document(
            title=path.name,
            source=str(path),
            document_type="word",
            content=content,
            metadata={
                "format": "docx",
                "paragraph_count": len(document.paragraphs),
                "table_count": len(document.tables),
            },
        )
