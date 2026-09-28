"""Reader for Microsoft Excel documents."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from models.document import Document
from readers.base_reader import BaseReader
from services.exceptions import DocumentReadError, InvalidDocumentError


class ExcelReader(BaseReader):
    """Read .xlsx files into a common Document model."""

    def can_read(self, source: str) -> bool:
        return Path(source).suffix.lower() == ".xlsx"

    def read(self, source: str) -> Document:
        path = Path(source)
        if not path.exists():
            raise InvalidDocumentError(
                "Unable to read this document.\n\nPlease verify that the file is valid and try again."
            )

        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:
            raise InvalidDocumentError(
                "Unable to read this document.\n\nPlease verify that the file is valid and try again."
            ) from exc

        sheet_sections: list[str] = []
        try:
            for worksheet in workbook.worksheets:
                rows: list[str] = []
                for row in worksheet.iter_rows(values_only=True):
                    values = ["" if value is None else str(value) for value in row]
                    rows.append(" | ".join(values))

                if rows:
                    sheet_sections.append(
                        f"=== Sheet: {worksheet.title} ===\n\n" + "\n".join(rows)
                    )
                else:
                    sheet_sections.append(f"=== Sheet: {worksheet.title} ===\n")
        except Exception as exc:
            raise DocumentReadError(
                "Unable to read this document.\n\nPlease verify that the file is valid and try again."
            ) from exc
        finally:
            workbook.close()

        content = "\n\n".join(sheet_sections).strip()

        return Document(
            title=path.name,
            source=str(path),
            document_type="excel",
            content=content,
            metadata={
                "format": "xlsx",
                "worksheet_count": len(workbook.worksheets),
                "worksheets": [worksheet.title for worksheet in workbook.worksheets],
            },
        )
