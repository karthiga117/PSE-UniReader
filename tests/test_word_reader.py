"""Tests for the Word reader."""

from __future__ import annotations

from pathlib import Path

from docx import Document as DocxDocument

from readers.word_reader import WordReader


def test_word_reader_paragraphs_and_headings(tmp_path: Path) -> None:
    file_path = tmp_path / "sample.docx"
    doc = DocxDocument()
    doc.add_heading("Title", level=1)
    doc.add_paragraph("Hello from Word")
    doc.save(file_path)

    result = WordReader().read(str(file_path))

    assert result.title == "sample.docx"
    assert result.document_type == "word"
    assert "Title" in result.content
    assert "Hello from Word" in result.content


def test_word_reader_table_extraction(tmp_path: Path) -> None:
    file_path = tmp_path / "table.docx"
    doc = DocxDocument()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Age"
    table.cell(1, 0).text = "John"
    table.cell(1, 1).text = "30"
    doc.save(file_path)

    result = WordReader().read(str(file_path))

    assert "Name" in result.content
    assert "John" in result.content
    assert "30" in result.content
