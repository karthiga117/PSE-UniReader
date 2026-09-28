"""Tests for the Excel reader."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from readers.excel_reader import ExcelReader


def test_excel_reader_multiple_worksheets_and_cells(tmp_path: Path) -> None:
    file_path = tmp_path / "sample.xlsx"
    workbook = Workbook()
    sheet_one = workbook.active
    sheet_one.title = "Employees"
    sheet_one["A1"] = "Name"
    sheet_one["B1"] = "Department"
    sheet_one["A2"] = "John"
    sheet_one["B2"] = "IT"

    sheet_two = workbook.create_sheet("Projects")
    sheet_two["A1"] = "Project"
    sheet_two["B1"] = "Status"
    sheet_two["A2"] = "Alpha"
    sheet_two["B2"] = "Active"

    workbook.save(file_path)
    workbook.close()

    result = ExcelReader().read(str(file_path))

    assert result.title == "sample.xlsx"
    assert result.document_type == "excel"
    assert "Employees" in result.content
    assert "John" in result.content
    assert "Alpha" in result.content
    assert result.metadata["worksheet_count"] == 2


def test_excel_reader_empty_worksheet(tmp_path: Path) -> None:
    file_path = tmp_path / "empty.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Empty"
    workbook.save(file_path)
    workbook.close()

    result = ExcelReader().read(str(file_path))

    assert "Empty" in result.content
    assert result.metadata["worksheet_count"] == 1
